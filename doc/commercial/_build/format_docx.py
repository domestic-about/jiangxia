"""Apply the review-document visual system after Pandoc generates the DOCX.

Adapted from chores/2026-08/T010-健康监测/_build/format_docx.py (V1.7 pipeline),
parametrized so every commercial document shares the same look.

Usage:
    python3 format_docx.py <docx> --footer "..." [--title ...] [--subject ...] [--author ...]
"""
import argparse
from typing import Optional

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

PAGE_BREAK_PREFIXES = ("附件", "附录", "目录")


def set_east_asia(font, name: str) -> None:
    font.name = name
    rpr = font._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), name)


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def style_runs(paragraph, size: float = 10.5, color: Optional[str] = None, bold: bool = False) -> None:
    for run in paragraph.runs:
        run.font.size = Pt(size)
        set_east_asia(run.font, "Microsoft YaHei")
        run.bold = bold or run.bold
        if color:
            run.font.color.rgb = RGBColor.from_string(color)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("docx")
    ap.add_argument("--footer", required=True)
    ap.add_argument("--title", default="")
    ap.add_argument("--subject", default="")
    ap.add_argument("--author", default="武汉市添达信息技术服务有限公司")
    args = ap.parse_args()

    doc = Document(args.docx)
    styles = doc.styles

    normal = styles["Normal"]
    normal.font.size = Pt(10.5)
    set_east_asia(normal.font, "Microsoft YaHei")
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.32

    for style_name, size, color in [
        ("Title", 25, "123B4A"),
        ("Subtitle", 13, "3E6978"),
        ("Heading 1", 16, "0F6078"),
        ("Heading 2", 13, "184E61"),
        ("Heading 3", 11.5, "246B86"),
    ]:
        if style_name not in styles:
            continue
        style = styles[style_name]
        style.font.size = Pt(size)
        style.font.bold = style_name != "Subtitle"
        style.font.color.rgb = RGBColor.from_string(color)
        set_east_asia(style.font, "Microsoft YaHei")
        style.paragraph_format.space_before = Pt(13 if style_name.startswith("Heading") else 4)
        style.paragraph_format.space_after = Pt(7)

    if "Caption" in styles:
        caption = styles["Caption"]
        caption.font.size = Pt(8.5)
        caption.font.color.rgb = RGBColor.from_string("6F8791")
        set_east_asia(caption.font, "Microsoft YaHei")
        caption.paragraph_format.space_before = Pt(1)
        caption.paragraph_format.space_after = Pt(4)

    for section in doc.sections:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Mm(21)
        section.bottom_margin = Mm(20)
        section.left_margin = Mm(22)
        section.right_margin = Mm(22)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.text = args.footer
        style_runs(footer, size=8.5, color="6F8791")

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        style_name = paragraph.style.name if paragraph.style is not None else ""
        if style_name == "Title":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if style_name == "Subtitle":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if style_name in ("Heading 1", "TOC Heading") and text.startswith(PAGE_BREAK_PREFIXES):
            paragraph.paragraph_format.page_break_before = True
        if text == "目录":
            paragraph.paragraph_format.page_break_before = True
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if paragraph._p.xpath(".//w:drawing"):
            paragraph.paragraph_format.keep_together = True

    for table in doc.tables:
        table_font_size = 9.5 if len(table.columns) >= 4 else 10.5
        for row_idx, row in enumerate(table.rows):
            tr_pr = row._tr.get_or_add_trPr()
            cant_split = OxmlElement("w:cantSplit")
            tr_pr.append(cant_split)
            for cell in row.cells:
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                cell.margin_top = Mm(1.2)
                cell.margin_bottom = Mm(1.2)
                cell.margin_left = Mm(1.4)
                cell.margin_right = Mm(1.4)
                if row_idx == 0:
                    shade(cell, "17647A")
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(1.5)
                    paragraph.paragraph_format.line_spacing = 1.15
                    style_runs(
                        paragraph,
                        size=table_font_size,
                        color="FFFFFF" if row_idx == 0 else None,
                        bold=row_idx == 0,
                    )

    if args.title:
        doc.core_properties.title = args.title
    if args.subject:
        doc.core_properties.subject = args.subject
    doc.core_properties.author = args.author
    doc.save(args.docx)


if __name__ == "__main__":
    main()
